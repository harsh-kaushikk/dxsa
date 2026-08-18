import re
from pathlib import Path
from xml.sax.saxutils import escape

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import Image as RLImage
from reportlab.platypus import PageBreak, Paragraph, Preformatted, SimpleDocTemplate, Spacer


INPUT_MD = Path("/workspace/7007SCN_Assignment_Submission_StrictOrdered.md")
OUT_DOCX = Path("/workspace/7007SCN_Assignment_Submission_StrictOrdered_FIXED.docx")
OUT_PDF = Path("/workspace/7007SCN_Assignment_Submission_StrictOrdered_FIXED.pdf")


def normalize_text(text: str) -> str:
    replacements = {
        "\u2013": "-",
        "\u2014": "-",
        "\u2018": "'",
        "\u2019": "'",
        "\u201c": '"',
        "\u201d": '"',
        "\u00a3": "GBP ",
    }
    for old, new in replacements.items():
        text = text.replace(old, new)
    return text


def shade_docx_paragraph(paragraph, fill="F3F3F3"):
    p_pr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    p_pr.append(shd)


def parse_blocks(md_text: str):
    blocks = []
    lines = md_text.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        if stripped == r"\newpage":
            blocks.append(("pagebreak", ""))
            i += 1
            continue

        if stripped.startswith("```"):
            lang = stripped[3:].strip()
            i += 1
            code_lines = []
            while i < len(lines) and not lines[i].strip().startswith("```"):
                code_lines.append(lines[i])
                i += 1
            if i < len(lines):
                i += 1
            blocks.append(("code", (lang, "\n".join(code_lines))))
            continue

        img_match = re.match(r"!\[(.*?)\]\((.*?)\)", stripped)
        if img_match:
            blocks.append(("image", img_match.group(2)))
            i += 1
            continue

        h_match = re.match(r"^(#{1,6})\s+(.*)$", stripped)
        if h_match:
            level = len(h_match.group(1))
            blocks.append(("heading", (level, h_match.group(2).strip())))
            i += 1
            continue

        if not stripped:
            i += 1
            continue

        para_lines = [line]
        i += 1
        while i < len(lines):
            nxt = lines[i].strip()
            if not nxt or nxt == r"\newpage" or nxt.startswith("```") or re.match(r"!\[.*\]\(.*\)", nxt) or re.match(r"^#{1,6}\s+", nxt):
                break
            para_lines.append(lines[i])
            i += 1
        text = " ".join(x.strip() for x in para_lines).strip()
        blocks.append(("paragraph", text))
    return blocks


def add_markdown_runs(paragraph, text):
    text = text.replace("  ", " ").strip()
    idx = 0
    for m in re.finditer(r"\*\*(.+?)\*", text):
        start, end = m.span()
        if start > idx:
            paragraph.add_run(text[idx:start])
        run = paragraph.add_run(m.group(1))
        run.bold = True
        idx = end
    if idx < len(text):
        paragraph.add_run(text[idx:])


def build_docx(blocks):
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(11)

    for block_type, value in blocks:
        if block_type == "pagebreak":
            doc.add_page_break()
        elif block_type == "heading":
            level, text = value
            p = doc.add_paragraph()
            run = p.add_run(normalize_text(text))
            run.bold = True
            run.font.name = "Times New Roman"
            run.font.size = Pt({1: 16, 2: 14, 3: 12}.get(level, 11))
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(3)
        elif block_type == "paragraph":
            p = doc.add_paragraph()
            add_markdown_runs(p, normalize_text(value))
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.line_spacing = 1.15
        elif block_type == "code":
            _, code_text = value
            for code_line in normalize_text(code_text).splitlines() or [""]:
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Inches(0.2)
                p.paragraph_format.right_indent = Inches(0.2)
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(0)
                shade_docx_paragraph(p)
                run = p.add_run(code_line)
                run.font.name = "Courier New"
                run.font.size = Pt(9.5)
        elif block_type == "image":
            img_path = (INPUT_MD.parent / value).resolve()
            if img_path.exists():
                doc.add_picture(str(img_path), width=Inches(6.2))

    doc.save(OUT_DOCX)


def inline_markdown_to_html(text: str) -> str:
    text = escape(normalize_text(text))
    text = re.sub(r"\*\*(.+?)\*\*", r"<b>\1</b>", text)
    text = re.sub(r"\*(.+?)\*", r"<i>\1</i>", text)
    return text


def build_pdf(blocks):
    pdf = SimpleDocTemplate(
        str(OUT_PDF),
        pagesize=A4,
        rightMargin=1.8 * cm,
        leftMargin=1.8 * cm,
        topMargin=1.6 * cm,
        bottomMargin=1.6 * cm,
    )
    styles = getSampleStyleSheet()
    normal = ParagraphStyle(
        "NormalTNR",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=10.5,
        leading=14,
        spaceAfter=5,
    )
    h1 = ParagraphStyle("H1", parent=styles["Heading1"], fontName="Helvetica-Bold", fontSize=15, leading=18, spaceBefore=8, spaceAfter=6)
    h2 = ParagraphStyle("H2", parent=styles["Heading2"], fontName="Helvetica-Bold", fontSize=13, leading=16, spaceBefore=6, spaceAfter=5)
    h3 = ParagraphStyle("H3", parent=styles["Heading3"], fontName="Helvetica-Bold", fontSize=11.5, leading=14, spaceBefore=5, spaceAfter=4)
    code_style = ParagraphStyle(
        "Code",
        fontName="Courier",
        fontSize=8.8,
        leading=11,
        leftIndent=8,
        rightIndent=8,
        backColor="#F3F3F3",
        spaceBefore=2,
        spaceAfter=2,
    )

    story = []
    for block_type, value in blocks:
        if block_type == "pagebreak":
            story.append(PageBreak())
        elif block_type == "heading":
            level, text = value
            style = h1 if level == 1 else h2 if level == 2 else h3 if level == 3 else normal
            story.append(Paragraph(inline_markdown_to_html(text), style))
        elif block_type == "paragraph":
            story.append(Paragraph(inline_markdown_to_html(value), normal))
        elif block_type == "code":
            _, code_text = value
            story.append(Preformatted(normalize_text(code_text), code_style, dedent=0))
            story.append(Spacer(1, 3))
        elif block_type == "image":
            img_path = (INPUT_MD.parent / value).resolve()
            if img_path.exists():
                img = RLImage(str(img_path))
                orig_w, orig_h = img.imageWidth, img.imageHeight
                target_w = 16.2 * cm
                img.drawWidth = target_w
                img.drawHeight = orig_h * (target_w / orig_w)
                story.append(img)
                story.append(Spacer(1, 6))

    pdf.build(story)


def main():
    md_text = INPUT_MD.read_text(encoding="utf-8")
    blocks = parse_blocks(md_text)
    build_docx(blocks)
    build_pdf(blocks)
    print(f"Created: {OUT_DOCX}")
    print(f"Created: {OUT_PDF}")


if __name__ == "__main__":
    main()
